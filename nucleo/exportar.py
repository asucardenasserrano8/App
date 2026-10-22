"""
Módulo generado a partir de la mejora acordada: exportar el reporte a PDF y
a Excel, con gráficas y formato profesional (no solo texto plano).
"""
from config import *

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
)
from reportlab.lib.enums import TA_CENTER

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

# Color principal de la marca (el mismo morado/índigo del degradado de la app)
_COLOR_MARCA = "#667eea"
_COLOR_MARCA_RGB = colors.HexColor(_COLOR_MARCA)


def _categoria_scoring(score):
    if score >= 70:
        return "Sólida", colors.HexColor("#155724")
    if score >= 50:
        return "Aceptable", colors.HexColor("#856404")
    return "Débil", colors.HexColor("#721c24")


def _grafica_precio_png(datos, nombre):
    """Genera una gráfica de precio como imagen PNG en memoria (para
    incrustarla en el PDF)."""
    fig, ax = plt.subplots(figsize=(6.5, 2.8), dpi=150)
    ax.plot(datos.index, datos["Close"], color="#667eea", linewidth=1.5)
    ax.set_title(f"Precio de cierre — {nombre}", fontsize=10)
    ax.set_ylabel("Precio ($)", fontsize=8)
    ax.tick_params(axis="both", labelsize=7)
    ax.grid(alpha=0.3)
    fig.tight_layout()
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png")
    plt.close(fig)
    buffer.seek(0)
    return buffer


def _sin_emoji(texto):
    """La fuente estándar del PDF no soporta emoji (salen como cuadros
    negros) — los quitamos y dejamos solo la palabra (ej. '🟢 Excelente'
    queda 'Excelente')."""
    for emoji in ("🟢", "🟡", "🔴", "✅", "❌", "⚠️"):
        texto = texto.replace(emoji, "")
    return texto.strip()


def generar_reporte_pdf(stonk, info, datos, scoring, metricas, nombre):
    """Arma un PDF con portada, tabla de métricas con formato, gráfica de
    precio y el detalle de scoring — en vez de solo texto plano."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=letter,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        leftMargin=2 * cm, rightMargin=2 * cm,
    )
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        "TituloIndice", parent=estilos["Title"], textColor=_COLOR_MARCA_RGB, fontSize=22,
    )
    estilo_subtitulo = ParagraphStyle(
        "SubtituloIndice", parent=estilos["Normal"], alignment=TA_CENTER,
        textColor=colors.grey, fontSize=10,
    )
    estilo_seccion = ParagraphStyle(
        "SeccionIndice", parent=estilos["Heading2"], textColor=_COLOR_MARCA_RGB,
        spaceBefore=14, spaceAfter=6,
    )

    elementos = []
    elementos.append(Paragraph(f"Índice — Reporte de {nombre}", estilo_titulo))
    elementos.append(Paragraph(f"Símbolo: {stonk}", estilo_subtitulo))
    elementos.append(Paragraph(
        f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}", estilo_subtitulo
    ))
    elementos.append(Spacer(1, 0.5 * cm))

    # --- Resumen ejecutivo (scoring) ---
    categoria, color_categoria = _categoria_scoring(scoring)
    elementos.append(Paragraph("Resumen Ejecutivo", estilo_seccion))
    tabla_resumen = Table(
        [["Puntuación General", "Categoría", "Precio Actual"],
         [f"{scoring}/100", categoria,
          f"${info.get('currentPrice', info.get('regularMarketPrice', 'N/D'))}"]],
        colWidths=[5.5 * cm, 5.5 * cm, 5.5 * cm],
    )
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _COLOR_MARCA_RGB),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("TEXTCOLOR", (1, 1), (1, 1), color_categoria),
        ("FONTNAME", (1, 1), (1, 1), "Helvetica-Bold"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    elementos.append(tabla_resumen)

    # --- Gráfica de precio ---
    if datos is not None and not datos.empty and "Close" in datos.columns:
        elementos.append(Paragraph("Evolución del Precio", estilo_seccion))
        imagen_buffer = _grafica_precio_png(datos, nombre)
        elementos.append(RLImage(imagen_buffer, width=16 * cm, height=6.5 * cm))

    # --- Tabla de métricas fundamentales (scoring detallado) ---
    if metricas:
        elementos.append(Paragraph("Métricas Fundamentales (Scoring)", estilo_seccion))
        filas = [["Métrica", "Evaluación"]] + [[k, _sin_emoji(v)] for k, v in metricas.items()]
        tabla_metricas = Table(filas, colWidths=[8 * cm, 8 * cm])
        tabla_metricas.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _COLOR_MARCA_RGB),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8f9fa")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ]))
        elementos.append(tabla_metricas)

    # --- Datos generales de la empresa ---
    elementos.append(Paragraph("Información General", estilo_seccion))
    datos_generales = [
        ["Sector", info.get("sector", "N/D")],
        ["Industria", info.get("industry", "N/D")],
        ["Market Cap", f"{info.get('marketCap', 'N/D'):,}" if isinstance(info.get("marketCap"), (int, float)) else "N/D"],
        ["País", info.get("country", "N/D")],
    ]
    tabla_generales = Table(datos_generales, colWidths=[5 * cm, 11 * cm])
    tabla_generales.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
    ]))
    elementos.append(tabla_generales)

    elementos.append(Spacer(1, 0.8 * cm))
    elementos.append(Paragraph(
        "Este reporte fue generado con fines educativos. No constituye asesoramiento financiero.",
        ParagraphStyle("Disclaimer", parent=estilos["Normal"], fontSize=8, textColor=colors.grey,
                       alignment=TA_CENTER),
    ))

    doc.build(elementos)
    buffer.seek(0)
    return buffer


def generar_reporte_excel(stonk, info, datos, scoring, metricas, nombre):
    """Arma un Excel con una hoja de resumen (con formato y colores) y una
    hoja de histórico de precios con una gráfica NATIVA de Excel (editable
    dentro de Excel, no una imagen pegada)."""
    wb = openpyxl.Workbook()

    # ---------- Hoja 1: Resumen ----------
    ws = wb.active
    ws.title = "Resumen"

    relleno_marca = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
    fuente_titulo = Font(size=16, bold=True, color="667EEA")
    fuente_encabezado = Font(bold=True, color="FFFFFF")
    borde_fino = Border(*(Side(style="thin", color="D0D0D0"),) * 4)

    ws["B2"] = f"Índice — Reporte de {nombre} ({stonk})"
    ws["B2"].font = fuente_titulo
    ws["B3"] = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}"
    ws["B3"].font = Font(italic=True, color="808080")

    categoria, _ = _categoria_scoring(scoring)
    ws["B5"] = "Puntuación General"
    ws["C5"] = "Categoría"
    ws["D5"] = "Precio Actual"
    for celda in ("B5", "C5", "D5"):
        ws[celda].fill = relleno_marca
        ws[celda].font = fuente_encabezado
        ws[celda].alignment = Alignment(horizontal="center")
    ws["B6"] = f"{scoring}/100"
    ws["C6"] = categoria
    ws["D6"] = info.get("currentPrice", info.get("regularMarketPrice", "N/D"))
    for celda in ("B6", "C6", "D6"):
        ws[celda].alignment = Alignment(horizontal="center")
        ws[celda].border = borde_fino

    fila = 8
    ws.cell(row=fila, column=2, value="Métrica").font = fuente_encabezado
    ws.cell(row=fila, column=3, value="Evaluación").font = fuente_encabezado
    for col in (2, 3):
        ws.cell(row=fila, column=col).fill = relleno_marca
    for nombre_metrica, evaluacion in (metricas or {}).items():
        fila += 1
        ws.cell(row=fila, column=2, value=nombre_metrica).border = borde_fino
        celda_valor = ws.cell(row=fila, column=3, value=evaluacion)
        celda_valor.border = borde_fino

    fila += 2
    info_general = [
        ("Sector", info.get("sector", "N/D")),
        ("Industria", info.get("industry", "N/D")),
        ("Market Cap", info.get("marketCap", "N/D")),
        ("País", info.get("country", "N/D")),
    ]
    for etiqueta, valor in info_general:
        ws.cell(row=fila, column=2, value=etiqueta).font = Font(bold=True)
        ws.cell(row=fila, column=3, value=valor)
        fila += 1

    for col, ancho in zip("ABCD", (3, 26, 20, 20)):
        ws.column_dimensions[col].width = ancho

    # ---------- Hoja 2: Histórico de precios + gráfica nativa ----------
    if datos is not None and not datos.empty and "Close" in datos.columns:
        ws2 = wb.create_sheet("Histórico de Precios")
        ws2.append(["Fecha", "Cierre"])
        ws2["A1"].font = fuente_encabezado
        ws2["B1"].font = fuente_encabezado
        ws2["A1"].fill = relleno_marca
        ws2["B1"].fill = relleno_marca

        for fecha, precio in zip(datos.index, datos["Close"]):
            try:
                ws2.append([fecha.strftime("%Y-%m-%d"), float(precio)])
            except Exception:
                continue

        n_filas = ws2.max_row
        grafica = LineChart()
        grafica.title = f"Precio de cierre — {nombre}"
        grafica.style = 2
        grafica.y_axis.title = "Precio ($)"
        grafica.x_axis.title = "Fecha"
        datos_ref = Reference(ws2, min_col=2, min_row=1, max_row=n_filas)
        categorias_ref = Reference(ws2, min_col=1, min_row=2, max_row=n_filas)
        grafica.add_data(datos_ref, titles_from_data=True)
        grafica.set_categories(categorias_ref)
        grafica.width = 24
        grafica.height = 10
        ws2.add_chart(grafica, f"D2")

        ws2.column_dimensions["A"].width = 14
        ws2.column_dimensions["B"].width = 12

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
